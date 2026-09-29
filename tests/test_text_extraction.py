"""Unit tests for text extraction and normalization service."""

from pathlib import Path
import tempfile
import unittest

import docx
from app.services.text_extraction import (
    extract_document_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_txt,
    normalize_text,
)

SAMPLE_PDF_BYTES = (
    b"%PDF-1.4\n"
    b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
    b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
    b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
    b"4 0 obj << /Length 55 >> stream\n"
    b"BT\n"
    b"/F1 18 Tf\n"
    b"0 0 Td\n"
    b"(Regulatory Affairs Guidance Text) Tj\n"
    b"ET\n"
    b"endstream endobj\n"
    b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
    b"xref\n"
    b"0 6\n"
    b"0000000000 65535 f\x20\n"
    b"0000000009 00000 n\x20\n"
    b"0000000058 00000 n\x20\n"
    b"0000000115 00000 n\x20\n"
    b"0000000246 00000 n\x20\n"
    b"0000000353 00000 n\x20\n"
    b"trailer << /Size 6 /Root 1 0 R >>\n"
    b"startxref\n"
    b"426\n"
    b"%%EOF"
)


class TestTextExtraction(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_extract_text_from_txt_utf8(self):
        txt_path = self.dir_path / "sample.txt"
        content = "Section 1: Regulatory Scope and Definitions.\nLine 2."
        txt_path.write_bytes(content.encode("utf-8"))

        extracted = extract_text_from_txt(txt_path)
        self.assertEqual(extracted, content)

    def test_extract_text_from_txt_latin1_fallback(self):
        txt_path = self.dir_path / "sample_latin1.txt"
        content = "Regulatory temperature: 25\xb0C."
        txt_path.write_bytes(content.encode("latin-1"))

        extracted = extract_text_from_txt(txt_path)
        self.assertIn("25", extracted)

    def test_extract_text_from_txt_file_not_found(self):
        missing = self.dir_path / "non_existent.txt"
        with self.assertRaises(FileNotFoundError):
            extract_text_from_txt(missing)

    def test_extract_text_from_docx_paragraphs_and_tables(self):
        docx_path = self.dir_path / "sample.docx"
        doc = docx.Document()
        doc.add_paragraph("Section 1.0 Regulatory Overview")
        doc.add_paragraph("Section 1.1 Requirements")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Active Substance"
        table.rows[0].cells[1].text = "Ibuprofen"
        doc.save(str(docx_path))

        extracted = extract_text_from_docx(docx_path)
        self.assertIn("Section 1.0 Regulatory Overview", extracted)
        self.assertIn("Section 1.1 Requirements", extracted)
        self.assertIn("Active Substance", extracted)
        self.assertIn("Ibuprofen", extracted)

    def test_extract_text_from_docx_file_not_found(self):
        missing = self.dir_path / "non_existent.docx"
        with self.assertRaises(FileNotFoundError):
            extract_text_from_docx(missing)

    def test_extract_text_from_pdf(self):
        pdf_path = self.dir_path / "sample.pdf"
        pdf_path.write_bytes(SAMPLE_PDF_BYTES)

        extracted = extract_text_from_pdf(pdf_path)
        self.assertIn("Regulatory Affairs Guidance Text", extracted)

    def test_extract_text_from_pdf_file_not_found(self):
        missing = self.dir_path / "non_existent.pdf"
        with self.assertRaises(FileNotFoundError):
            extract_text_from_pdf(missing)

    def test_normalize_text_whitespace_and_newlines(self):
        raw = "\r\n  Line 1   with   spaces.  \r\n\r\n\r\n\r\nLine 2.\r\n"
        normalized = normalize_text(raw)
        expected = "Line 1 with spaces.\n\nLine 2."
        self.assertEqual(normalized, expected)

    def test_normalize_text_unicode_spaces_and_control_chars(self):
        raw = "Word1\u00a0Word2\ufeff\u200b\x07"
        normalized = normalize_text(raw)
        self.assertEqual(normalized, "Word1 Word2")

    def test_normalize_text_empty(self):
        self.assertEqual(normalize_text(""), "")
        self.assertEqual(normalize_text("   \t\r\n  "), "")

    def test_extract_document_text_dispatcher(self):
        txt_path = self.dir_path / "test_doc.txt"
        txt_path.write_bytes(b"Paragraph one.\r\n\r\nParagraph two.")

        result = extract_document_text(txt_path, "txt")
        self.assertEqual(result, "Paragraph one.\n\nParagraph two.")

        result2 = extract_document_text(txt_path, ".txt")
        self.assertEqual(result2, "Paragraph one.\n\nParagraph two.")

    def test_extract_document_text_unsupported_type(self):
        txt_path = self.dir_path / "test_doc.txt"
        txt_path.write_bytes(b"Content")

        with self.assertRaises(ValueError):
            extract_document_text(txt_path, "xyz")


if __name__ == "__main__":
    unittest.main()
